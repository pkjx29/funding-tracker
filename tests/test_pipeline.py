"""End-to-end on the made-up records in conftest.py: match, reconcile, check, write the workbook."""
import re
import zipfile
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from funding_tracker.checks import CHECKS, run_checks
from funding_tracker.excel import write_workbook
from funding_tracker.fx import inr_crore_to_usd_mn, rate_on
from funding_tracker.pipeline import coverage_table
from funding_tracker.reconcile import build_rounds
from funding_tracker.report import headline, stats_block

TODAY = date(2026, 10, 11)


@pytest.fixture
def built(records, rates):
    rounds, matched, merge_log, unmatched = build_rounds(records, rates)
    exceptions, summary = run_checks(rounds, matched, rates, TODAY)
    rounds["flags"] = rounds["round_id"].map(
        exceptions.groupby("round_id")["check_id"].agg(lambda ids: " ".join(sorted(set(ids))))).fillna("")
    return {"rounds": rounds, "records": matched, "merge_log": merge_log, "unmatched_headlines": unmatched,
            "exceptions": exceptions, "checks_summary": summary,
            "coverage": coverage_table(rounds, pd.DataFrame(columns=["source"]))}


def round_of(built, company):
    rounds = built["rounds"]
    return rounds[rounds["company"] == company].iloc[0]


def flags(built, company):
    return set(round_of(built, company)["flags"].split())


def test_fx_helpers(rates):
    assert rate_on(rates, "2026-10-08") == 96.0
    assert rate_on(rates, "2031-01-01") == 96.0                     # clamped to the last rate held
    assert inr_crore_to_usd_mn(96.0, 96.0) == pytest.approx(10.0)   # Rs 96 crore = $10 million at 96


def test_reports_are_grouped_into_rounds(built):
    rounds = built["rounds"]
    assert len(rounds) == 7                                          # A, B, C twice, D, E, F
    assert sorted(rounds["company"]) == ["Acme Goods", "Bolt Pay", "Cargo Lane", "Cargo Lane", "Dash Labs",
                                         "Mosaic Wellness", "Yonder Labs"]
    assert list(built["unmatched_headlines"]["record_id"]) == ["ET-3"]
    assert (built["records"].set_index("record_id").loc["ENT-5", "round_id"]) == ""


def test_source_of_truth_rules(built):
    acme = round_of(built, "Acme Goods")
    assert acme["sources"] == "entrackr | inc42 | et | vccircle" and acme["report_count"] == 4
    assert acme["announced_date"] == "2026-10-08"                    # the earliest report
    assert (acme["amount_source"], acme["amount_as_reported"], acme["amount_inr_cr"]) == ("entrackr", "Rs 332 crore", 332.0)
    assert acme["amount_usd_mn"] == pytest.approx(332 * 10 / 96, abs=0.001)   # converted at the reference rate
    assert acme["stage"] == "Series C" and acme["stage_group"] == "Growth"
    assert acme["lead_investors"] == "Northstar Capital Partners"   # one investor, named once
    assert acme["investors"] == "Northstar Capital Partners; Lake VC"
    assert (acme["sector"], acme["city"]) == ("Ecommerce", "Delhi NCR")

    mosaic = round_of(built, "Mosaic Wellness")                      # the majority spelling, not the typo
    assert mosaic["name_variants"] == "Mosiac Wellness | Mosaic Wellness"
    assert mosaic["sector"] == "Healthtech"                          # "Health Tech" is an alias
    yonder = round_of(built, "Yonder Labs")                          # created by an ET article alone
    assert (yonder["sources"], yonder["city"], yonder["sector"]) == ("et", "Chennai", "Advanced Hardware & Technology")
    dash = round_of(built, "Dash Labs")
    assert (dash["sector"], dash["sector_basis"]) == ("Artificial Intelligence", "keyword rule on description")
    assert pd.isna(dash["amount_usd_mn"]) and dash["city"] == "Unknown"


def test_merge_log_says_how_each_duplicate_was_recognised(built):
    basis = built["merge_log"].set_index("record_id")["match_basis"]
    assert basis["INC-20261010-01"] == "spacing or punctuation variant"
    assert basis["ET-1"] == "identical name"
    assert "fuzzy" in basis["ENT-4"]
    assert len(built["merge_log"]) == 6                              # 13 matched reports -> 7 rounds


def test_checks_flag_the_right_rounds(built):
    assert len(CHECKS) == 18 and len(built["checks_summary"]) == 18
    assert flags(built, "Acme Goods") == set()
    assert {"C14", "C15", "C17", "C18"} <= flags(built, "Bolt Pay")   # scale slip, amount, lead and stage disagree
    assert {"C03", "C04", "C05", "C07"} <= flags(built, "Dash Labs")  # the thin report
    later_cargo = built["rounds"][built["rounds"]["company"] == "Cargo Lane"].iloc[1]
    assert "C09" in later_cargo["flags"].split()                     # second tranche within 45 days
    counts = built["checks_summary"].set_index("check_id")["rounds_flagged"]
    assert counts[["C01", "C02", "C08", "C10", "C12", "C13"]].sum() == 0


def test_inr_usd_mismatch_is_caught(records, rates):
    records.loc[records["record_id"] == "ENT-1", "amount_usd_mn"] = 350.0    # Rs 332 crore is not $350 million
    rounds, matched, _, _ = build_rounds(records, rates)
    exceptions, _ = run_checks(rounds, matched, rates, TODAY)
    assert "ENT-1" in exceptions.loc[exceptions["check_id"] == "C13", "detail"].iloc[0]


def test_headline_numbers(built):
    numbers = headline(built)
    assert (numbers["rounds"], numbers["reports_matched"], numbers["duplicates_merged"]) == (7, 13, 6)
    assert numbers["variant_merges"] == 2 and numbers["scale_mismatches"] == 1   # AcmeGoods spacing, Mosiac typo
    assert "Funding rounds tracked | **7**" in stats_block(built)


def test_workbook_has_pivot_tables_tables_and_stored_formula_results(built, rates, tmp_path):
    path = tmp_path / "tracker.xlsx"
    write_workbook(built, path, rates)
    workbook = load_workbook(path)
    assert workbook.sheetnames[:5] == ["Dashboard", "Rounds", "By Sector", "By Stage", "By City"]
    for sheet, field in (("By Sector", "Sector"), ("By Stage", "Stage"), ("By City", "City")):
        (pivot,) = workbook[sheet]._pivots
        assert pivot.rowHeaderCaption == field
        assert [f.name for f in pivot.dataFields] == ["Rounds", "Total USD mn", "Average USD mn"]
        assert pivot.cache.cacheSource.worksheetSource.name == "Rounds" and pivot.cache.refreshOnLoad
    assert "Rounds" in workbook["Rounds"].tables and workbook["Rounds"].max_row == 8
    assert len(workbook["Dashboard"]._charts) == 5
    assert workbook["Dashboard"]["A5"].value.startswith("=COUNTA(")
    with zipfile.ZipFile(path) as archive:                           # the first sheet is the dashboard
        dashboard = archive.read("xl/worksheets/sheet1.xml").decode()
    assert '<f>COUNTA(Rounds!$A$2:$A$8)</f><v>7</v>' in dashboard
    assert not re.search(r'<c r="\w+"(?: s="\d+")?><f>[^<]*</f><v></v>', dashboard)   # no number left without its result
