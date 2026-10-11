"""Write the Excel workbook: dashboard, data, PivotTables, checks and rules.

The PivotTables are real ones. They are defined on the ``Rounds`` table and
refresh when the file is opened; the cells underneath already hold the same
numbers, so the workbook also reads correctly in viewers that cannot refresh.
"""
from __future__ import annotations

import re
import shutil
import zipfile
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.fields import Index
from openpyxl.pivot.fields import Text as SharedText
from openpyxl.pivot.table import (DataField, FieldItem, Location, PivotField, PivotTableStyle, RowColField,
                                  RowColItem, TableDefinition)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import parse as P
from .checks import CHECKS
from .config import AMOUNT_TOLERANCE, SOURCE_LABELS, START_DATE
from .reconcile import SOURCE_RULES

# Colours: one blue for every chart (each chart shows a single series), neutral ink for all text.
INK, MUTED, FAINT = "0B0B0B", "52514E", "8A8985"
GRID, PANEL, WHITE = "E6E5E1", "F6F5F2", "FFFFFF"
BLUE = "2A78D6"
SEVERITY_COLOUR = {"error": "D03B3B", "warning": "FAB219", "info": "8A8985"}

TABLE_STYLE = "TableStyleMedium2"
PIVOT_STYLE = "PivotStyleLight16"
DATE_FORMAT, MONTH_FORMAT = "dd-mmm-yyyy", "mmm-yy"

# Rounds sheet: (header, column in rounds.csv, width, number format)
ROUND_COLUMNS = [
    ("Round ID", "round_id", 24, None),
    ("Company", "company", 26, None),
    ("Announced", "announced_date", 13, DATE_FORMAT),
    ("Month", "month", 9, MONTH_FORMAT),
    ("Amount (USD mn)", "amount_usd_mn", 16, "#,##0.00"),
    ("Amount (INR cr)", "amount_inr_cr", 16, "#,##0.00"),
    ("As reported", "amount_as_reported", 16, None),
    ("Stage", "stage", 14, None),
    ("Stage group", "stage_group", 13, None),
    ("Sector", "sector", 26, None),
    ("City", "city", 16, None),
    ("Lead investors", "lead_investors", 34, None),
    ("Investors", "investors", 44, None),
    ("Sources", "sources", 26, None),
    ("Source count", "source_count", 13, "0"),
    ("Amount spread %", "amount_spread_pct", 16, "0.0"),
    ("Flags", "flags", 20, None),
    ("First source link", "first_url", 40, None),
]
HEADERS = [header for header, *_ in ROUND_COLUMNS]

def _rich(size: int, colour: str = MUTED, bold: bool = False) -> RichText:
    """Text properties for a chart axis."""
    properties = CharacterProperties(sz=size * 100, b=bold, solidFill=colour)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=properties), endParaRPr=properties)])


def _chart_title(text: str) -> Title:
    properties = CharacterProperties(sz=1100, b=True, solidFill=INK)
    paragraph = Paragraph(pPr=ParagraphProperties(defRPr=properties), r=[RegularTextRun(rPr=properties, t=text)])
    return Title(tx=Text(rich=RichText(p=[paragraph])), overlay=False)


def _hairline(colour: str = GRID) -> GraphicalProperties:
    return GraphicalProperties(ln=LineProperties(solidFill=colour, w=6350))


def _style_chart(chart, title: str, width: float, height: float) -> None:
    chart.title = _chart_title(title)
    chart.legend = None                                   # one series: the title already names it
    chart.width, chart.height = width, height
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.txPr = _rich(9)
        axis.spPr = _hairline()
        axis.title = None


def _bar_chart(data_ws, title: str, label_col: int, value_col: int, first: int, last: int,
               width: float, height: float):
    """Horizontal bars, largest at the top, value at the tip of each bar."""
    chart = BarChart()
    chart.type = "bar"
    chart.add_data(Reference(data_ws, min_col=value_col, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(data_ws, min_col=label_col, min_row=first, max_row=last))
    _style_chart(chart, title, width, height)
    chart.gapWidth = 55
    series = chart.series[0]
    series.graphicalProperties = GraphicalProperties(solidFill=BLUE, ln=LineProperties(noFill=True))
    chart.x_axis.scaling.orientation = "maxMin"           # first row at the top
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.y_axis.delete = True                            # values sit on the bars instead
    chart.y_axis.majorGridlines = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(chart.dataLabels, flag, False)
    chart.dataLabels.txPr = _rich(9, INK)
    chart.dataLabels.numFmt = "#,##0"
    return chart


def _time_chart(kind: str, data_ws, title: str, label_col: int, value_col: int, first: int, last: int,
                width: float, height: float, number_format: str):
    """Monthly columns or a monthly line, with a quiet value axis."""
    chart = LineChart() if kind == "line" else BarChart()
    if kind != "line":
        chart.type = "col"
        chart.gapWidth = 45
    chart.add_data(Reference(data_ws, min_col=value_col, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(data_ws, min_col=label_col, min_row=first, max_row=last))
    _style_chart(chart, title, width, height)
    series = chart.series[0]
    if kind == "line":
        series.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=BLUE, w=25400))
        series.smooth = False
        series.marker.symbol = "circle"
        series.marker.size = 5
        series.marker.graphicalProperties = GraphicalProperties(solidFill=BLUE, ln=LineProperties(solidFill=WHITE))
    else:
        series.graphicalProperties = GraphicalProperties(solidFill=BLUE, ln=LineProperties(noFill=True))
    chart.y_axis.majorGridlines = ChartLines(spPr=_hairline())
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.y_axis.number_format = number_format
    chart.x_axis.tickLblSkip = 3
    chart.x_axis.tickMarkSkip = 3
    return chart


# ------------------------------------------------------------------- tables
def _clean(value):
    """A value openpyxl can write: no NaN, no pandas types."""
    if value is None or (isinstance(value, float) and pd.isna(value)) or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if hasattr(value, "item"):
        return value.item()
    return value


def _write_table(ws, frame: pd.DataFrame, name: str, widths: list[float], formats: list[str | None],
                 top: int = 1) -> int:
    """Write ``frame`` as an Excel table starting at row ``top``. Returns the last row."""
    for col, header in enumerate(frame.columns, 1):
        ws.cell(row=top, column=col, value=header)
    for offset, values in enumerate(frame.itertuples(index=False), 1):
        for col, value in enumerate(values, 1):
            ws.cell(row=top + offset, column=col, value=_clean(value))
    last = top + max(len(frame), 1)
    for col, (width, number_format) in enumerate(zip(widths, formats), 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        if number_format:
            for row in range(top + 1, last + 1):
                ws.cell(row=row, column=col).number_format = number_format
    table = Table(displayName=name, ref=f"A{top}:{get_column_letter(len(frame.columns))}{last}")
    table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
    ws.add_table(table)
    return last


def _heading(ws, title: str, note: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color=INK)
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(size=10, color=MUTED)


def rounds_frame(rounds: pd.DataFrame) -> pd.DataFrame:
    """The Rounds sheet: friendly headers, real dates, one link per round."""
    frame = rounds.copy()
    frame["announced_date"] = pd.to_datetime(frame["announced_date"]).dt.date
    frame["month"] = frame["announced_date"].map(lambda day: day.replace(day=1))
    frame["first_url"] = frame["source_urls"].str.split(" | ", regex=False).str[0]
    frame = frame[[column for _, column, *_ in ROUND_COLUMNS]]
    frame.columns = HEADERS
    return frame


# -------------------------------------------------------------- PivotTables
def _pivot_cache(frame: pd.DataFrame, orders: dict[str, list[str]]) -> CacheDefinition:
    """One cache for all PivotTables, pointing at the Rounds table.

    No records are stored (``saveData=False``); Excel reads them from the table
    when the file opens. Row fields list their items in the order they should
    be displayed.
    """
    fields = []
    for header in frame.columns:
        values = frame[header]
        if header in orders:
            shared = SharedItems(_fields=[SharedText(v=item) for item in orders[header]], count=len(orders[header]))
        elif pd.api.types.is_numeric_dtype(values) and values.notna().any():
            shared = SharedItems(containsSemiMixedTypes=False, containsString=False, containsNumber=True,
                                 containsBlank=True if values.isna().any() else None,
                                 minValue=float(values.min()), maxValue=float(values.max()))
        else:
            shared = SharedItems()
        fields.append(CacheField(name=header, numFmtId=0, sharedItems=shared))
    source = CacheSource(type="worksheet", worksheetSource=WorksheetSource(name="Rounds"))
    return CacheDefinition(cacheSource=source, cacheFields=fields, refreshOnLoad=True, saveData=False,
                           refreshedBy="funding-tracker", createdVersion=6, refreshedVersion=6,
                           minRefreshableVersion=3)


def _add_pivot(ws, frame: pd.DataFrame, cache: CacheDefinition, row_field: str, order: list[str],
               name: str, top: int = 4) -> None:
    """Rounds, total and average amount by ``row_field``, as a PivotTable anchored at A{top}."""
    amount, round_id = "Amount (USD mn)", "Round ID"
    grouped = frame.groupby(row_field)[amount].agg(["size", "sum", "mean", "count"]).reindex(order)

    # Cells exactly as Excel lays the PivotTable out: a "Values" row, the field names, items, grand total.
    ws.cell(row=top, column=2, value="Values")
    for col, text in enumerate([row_field, "Rounds", "Total USD mn", "Average USD mn"], 1):
        ws.cell(row=top + 1, column=col, value=text).font = Font(bold=True)
    body = [(item, int(r["size"]), float(r["sum"]) if r["count"] else None,
             float(r["mean"]) if r["count"] else None) for item, r in grouped.iterrows()]
    body.append(("Grand Total", len(frame), float(frame[amount].sum()), float(frame[amount].mean())))
    for offset, values in enumerate(body, 2):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=top + offset, column=col, value=value)
            cell.number_format = {2: "#,##0", 3: "#,##0", 4: "#,##0.00"}.get(col, "General")
    end = top + 1 + len(body)
    for cell in ws[end]:
        cell.font = Font(bold=True)

    columns = list(frame.columns)
    pivot_fields = []
    for header in columns:
        if header == row_field:
            items = [FieldItem(x=i) for i in range(len(order))] + [FieldItem(t="default")]
            pivot_fields.append(PivotField(axis="axisRow", showAll=False, items=items))
        else:
            pivot_fields.append(PivotField(dataField=True if header in (amount, round_id) else None, showAll=False))
    pivot = TableDefinition(
        name=name, cacheId=1, dataCaption="Values", updatedVersion=6, minRefreshableVersion=3, createdVersion=6,
        useAutoFormatting=True, itemPrintTitles=True, indent=0, outline=True, outlineData=True,
        multipleFieldFilters=False, applyWidthHeightFormats=True, rowHeaderCaption=row_field,
        location=Location(ref=f"A{top}:D{end}", firstHeaderRow=1, firstDataRow=2, firstDataCol=1),
        pivotFields=pivot_fields,
        rowFields=[RowColField(x=columns.index(row_field))],
        rowItems=[RowColItem(x=[Index(v=i)]) for i in range(len(order))] + [RowColItem(t="grand", x=[Index(v=0)])],
        colFields=[RowColField(x=-2)],
        colItems=[RowColItem(x=[Index(v=0)]), RowColItem(i=1, x=[Index(v=1)]), RowColItem(i=2, x=[Index(v=2)])],
        dataFields=[
            DataField(name="Rounds", fld=columns.index(round_id), subtotal="count", baseField=0, baseItem=0, numFmtId=3),
            DataField(name="Total USD mn", fld=columns.index(amount), baseField=0, baseItem=0, numFmtId=3),
            DataField(name="Average USD mn", fld=columns.index(amount), subtotal="average", baseField=0, baseItem=0,
                      numFmtId=4),
        ],
        pivotTableStyleInfo=PivotTableStyle(name=PIVOT_STYLE, showRowHeaders=True, showColHeaders=True,
                                            showRowStripes=False, showColStripes=False, showLastColumn=True))
    pivot.cache = cache
    ws._pivots.append(pivot)
    for letter, width in zip("ABCD", (32, 12, 16, 18)):
        ws.column_dimensions[letter].width = width


def ranked(values: pd.Series) -> pd.Series:
    """Counts, largest first, with ties in alphabetical order so that every run gives the same order."""
    return values.value_counts().sort_index().sort_values(ascending=False, kind="stable")


def display_orders(frame: pd.DataFrame) -> dict[str, list[str]]:
    """Item order for each PivotTable: stages along the funding ladder, the rest by number of rounds."""
    by_count = lambda column: ranked(frame[column]).index.tolist()             # noqa: E731
    last = lambda items, tail: [i for i in items if i not in tail] + [t for t in tail if t in items]  # noqa: E731
    present = set(frame["Stage"])
    return {
        "Sector": last(by_count("Sector"), ["Unclassified"]),
        "Stage": [stage for stage in P.STAGES if stage in present],
        "City": last(by_count("City"), ["Outside India", "Unknown"]),
    }


# ---------------------------------------------------------------- dashboard
def _kpi(ws, column: int, label: str, formula: str, note: str, number_format: str) -> str:
    """One stat tile, three columns wide. Returns the address of its value cell."""
    left, right = get_column_letter(column), get_column_letter(column + 2)
    for row, height in ((4, 18), (5, 34), (6, 18)):
        ws.merge_cells(f"{left}{row}:{right}{row}")
        ws.row_dimensions[row].height = height
        for col in range(column, column + 3):
            ws.cell(row=row, column=col).fill = PatternFill("solid", start_color=PANEL)
    ws[f"{left}4"] = label
    ws[f"{left}4"].font = Font(size=10, color=MUTED)
    ws[f"{left}5"] = formula
    ws[f"{left}5"].font = Font(size=22, bold=True, color=INK)
    ws[f"{left}5"].number_format = number_format
    ws[f"{left}6"] = note
    ws[f"{left}6"].font = Font(size=9, color=FAINT)
    for row in (4, 5, 6):
        ws[f"{left}{row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return f"{left}5"


def _dashboard(wb: Workbook, frame: pd.DataFrame, tables: dict[str, pd.DataFrame], n_reports: int,
               cached: dict[str, dict[str, object]]) -> None:
    ws = wb["Dashboard"]
    ws.sheet_view.showGridLines = False
    for col in range(1, 19):
        ws.column_dimensions[get_column_letter(col)].width = 12
    n = len(frame) + 1                                    # last data row on the Rounds sheet
    column = {header: get_column_letter(i) for i, header in enumerate(HEADERS, 1)}
    rng = lambda header: f"Rounds!${column[header]}$2:${column[header]}${n}"     # noqa: E731
    amounts = frame["Amount (USD mn)"]
    first_day, last_day = frame["Announced"].min(), frame["Announced"].max()

    ws["A1"] = "India Startup Funding Tracker"
    ws["A1"].font = Font(size=18, bold=True, color=INK)
    ws["A2"] = (f"Rounds announced {first_day:%d %b %Y} to {last_day:%d %b %Y}, cross-checked across "
                f"{', '.join(SOURCE_LABELS.values())}. Built {date.today():%d %b %Y}.")
    ws["A2"].font = Font(size=10, color=MUTED)

    merge_log = tables["merge_log"]
    variants = int((merge_log["match_basis"] != "identical name").sum())
    multi = int((frame["Source count"] >= 2).sum())
    tolerance = AMOUNT_TOLERANCE * 100
    disagree = int((frame["Amount spread %"] > tolerance).sum())
    values = cached.setdefault("Dashboard", {})
    tiles = [
        ("Rounds tracked", f"=COUNTA({rng('Round ID')})", f"{multi / len(frame):.0%} confirmed by 2+ sources",
         "#,##0", len(frame)),
        ("Disclosed funding, USD bn", f"=SUM({rng('Amount (USD mn)')})/1000",
         f"{int(amounts.notna().sum()):,} rounds disclosed an amount", "0.00", float(amounts.sum()) / 1000),
        ("Median round, USD mn", f"=MEDIAN({rng('Amount (USD mn)')})", "half of disclosed rounds are smaller",
         "0.00", float(amounts.median())),
        ("Source reports matched", f"=COUNTA(Reports!$A$5:$A${n_reports + 4})", "articles, table rows and headlines",
         "#,##0", n_reports),
        ("Duplicate reports merged", "=J5-A5", f"{variants:,} needed name matching", "#,##0", n_reports - len(frame)),
        (f"Amount disagreements over {tolerance:.0f}%", f'=COUNTIF({rng("Amount spread %")},">{tolerance:g}")',
         "rounds where sources differ", "#,##0", disagree),
    ]
    for i, (label, formula, note, number_format, value) in enumerate(tiles):
        values[_kpi(ws, 1 + 3 * i, label, formula, note, number_format)] = value

    data = wb["Chart data"]
    months, stages, sectors, cities = (cached["_extent"][key] for key in ("months", "stages", "sectors", "cities"))
    ws.add_chart(_time_chart("col", data, "Rounds announced per month", 1, 3, 2, months + 1, 19.5, 7.5, "#,##0"), "A8")
    ws.add_chart(_time_chart("line", data, "Disclosed funding per month, USD mn", 1, 4, 2, months + 1, 19.5, 7.5,
                             "#,##0"), "J8")
    ws["A23"] = f"The last month shown is partial: data runs to {last_day:%d %b %Y}."
    ws["A23"].font = Font(size=9, color=FAINT)
    ws.add_chart(_bar_chart(data, "Rounds by stage", 6, 7, 2, stages + 1, 12.8, 10.5), "A24")
    ws.add_chart(_bar_chart(data, "Rounds by sector", 9, 10, 2, sectors + 1, 12.8, 10.5), "G24")
    ws.add_chart(_bar_chart(data, "Rounds by city, where a report states it", 12, 13, 2, cities + 1, 12.8, 10.5), "M24")

    # ---- look up one round by its ID (INDEX/MATCH on the Rounds table)
    top = 46
    ws[f"A{top}"] = "Look up a round"
    ws[f"A{top}"].font = Font(size=12, bold=True, color=INK)
    ws[f"A{top + 1}"] = "Round ID"
    ws[f"C{top + 1}"] = frame["Round ID"].iloc[-1]
    ws[f"C{top + 1}"].fill = PatternFill("solid", start_color="FFF7D6")
    ws[f"C{top + 1}"].border = Border(*(Side(style="thin", color=GRID),) * 4)
    ws.merge_cells(f"C{top + 1}:E{top + 1}")
    ws[f"F{top + 1}"] = "pick an ID from the list, or paste one from the Rounds sheet"
    ws[f"F{top + 1}"].font = Font(size=9, color=FAINT)
    choices = DataValidation(type="list", formula1=f"={rng('Round ID')}", allow_blank=False)
    ws.add_data_validation(choices)
    choices.add(f"C{top + 1}")
    chosen = frame.iloc[-1]
    shown = ["Company", "Announced", "Amount (USD mn)", "As reported", "Stage", "Sector", "City", "Lead investors",
             "Sources", "Flags"]
    for offset, header in enumerate(shown, 2):
        row = top + offset
        ws[f"A{row}"] = header
        ws[f"A{row}"].font = Font(color=MUTED)
        lookup = f"INDEX({rng(header)},MATCH($C${top + 1},{rng('Round ID')},0))"
        ws[f"C{row}"] = f'=IFERROR(IF({lookup}="","",{lookup}),"not found")'
        ws[f"C{row}"].alignment = Alignment(horizontal="left")
        ws[f"C{row}"].number_format = {"Announced": DATE_FORMAT, "Amount (USD mn)": "#,##0.00"}.get(header, "General")
        value = _clean(chosen[header])
        values[f"C{row}"] = "" if value is None else value
    ws[f"A{top + len(shown) + 3}"] = ("Flags are check IDs from the Checks sheet. PivotTables are on the By Sector, "
                                      "By Stage and By City sheets and refresh when the file opens.")
    ws[f"A{top + len(shown) + 3}"].font = Font(size=9, color=FAINT)
    ws.freeze_panes = "A4"


def _chart_data(wb: Workbook, frame: pd.DataFrame, orders: dict[str, list[str]],
                cached: dict[str, dict[str, object]]) -> None:
    """Small summary tables that feed the dashboard charts, written as live formulas."""
    ws = wb["Chart data"]
    n = len(frame) + 1
    column = {header: get_column_letter(i) for i, header in enumerate(HEADERS, 1)}
    rng = lambda header: f"Rounds!${column[header]}$2:${column[header]}${n}"     # noqa: E731
    values = cached.setdefault("Chart data", {})

    months = pd.date_range(frame["Month"].min(), frame["Month"].max(), freq="MS").date
    per_month = frame.groupby("Month")["Amount (USD mn)"].agg(["size", "sum"])
    for col, header in enumerate(["Month", "Month start", "Rounds", "Funding USD mn"], 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    for row, month in enumerate(months, 2):
        ws.cell(row=row, column=1, value=f"{month:%b %y}")
        ws.cell(row=row, column=2, value=month).number_format = DATE_FORMAT
        ws.cell(row=row, column=3, value=f"=COUNTIF({rng('Month')},B{row})")
        ws.cell(row=row, column=4, value=f"=SUMIF({rng('Month')},B{row},{rng('Amount (USD mn)')})")
        ws.cell(row=row, column=4).number_format = "#,##0.0"
        values[f"C{row}"] = int(per_month["size"].get(month, 0))
        values[f"D{row}"] = float(per_month["sum"].get(month, 0.0))

    city_counts = ranked(frame.loc[~frame["City"].isin(["Unknown", "Outside India"]), "City"]).head(12)
    blocks = [("Stage", [s for s in orders["Stage"]], 6),
              ("Sector", [s for s in orders["Sector"]][:14], 9),
              ("City", city_counts.index.tolist(), 12)]
    counts = {header: frame[header].value_counts() for header, _, _ in blocks}
    for header, items, col in blocks:
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
        ws.cell(row=1, column=col + 1, value="Rounds").font = Font(bold=True)
        letter = get_column_letter(col)
        for row, item in enumerate(items, 2):
            ws.cell(row=row, column=col, value=item)
            ws.cell(row=row, column=col + 1, value=f"=COUNTIF({rng(header)},{letter}{row})")
            values[f"{get_column_letter(col + 1)}{row}"] = int(counts[header].get(item, 0))
    for letter, width in zip("ABCDEFGHIJKLM", (10, 13, 9, 15, 3, 16, 9, 3, 30, 9, 3, 18, 9)):
        ws.column_dimensions[letter].width = width
    ws["O1"] = "These tables feed the Dashboard charts. They are formulas on the Rounds sheet, so they follow any edit."
    ws["O1"].font = Font(size=9, color=FAINT)
    cached["_extent"] = {"months": len(months), "stages": len(blocks[0][1]), "sectors": len(blocks[1][1]),
                         "cities": len(blocks[2][1])}


# -------------------------------------------------------------- other sheets
def _checks_sheet(wb: Workbook, summary: pd.DataFrame) -> None:
    ws = wb["Checks"]
    _heading(ws, f"The {len(CHECKS)} automated checks",
             "Every round is tested on each run. A flag sends the round to a person; nothing is corrected automatically.")
    frame = summary.rename(columns={"check_id": "ID", "category": "Category", "check": "Check", "rule": "Rule",
                                    "severity": "Severity", "rounds_checked": "Rounds checked",
                                    "rounds_flagged": "Rounds flagged", "share_flagged": "Share flagged"})
    last = _write_table(ws, frame, "Checks", [6, 14, 36, 78, 11, 15, 15, 14],
                        [None, None, None, None, None, "#,##0", "#,##0", "0.0%"], top=4)
    for row in range(5, last + 1):                         # a coloured mark beside the severity label
        severity = ws.cell(row=row, column=5).value
        ws.cell(row=row, column=9, value="●").font = Font(color=SEVERITY_COLOUR.get(severity, FAINT))
    ws.column_dimensions["I"].width = 3
    ws.freeze_panes = "A5"


def _plain_sheet(wb: Workbook, name: str, title: str, note: str, frame: pd.DataFrame, table: str,
                 widths: list[float], formats: list[str | None]) -> None:
    ws = wb[name]
    _heading(ws, title, note)
    _write_table(ws, frame, table, widths, formats, top=4)
    ws.freeze_panes = "A5"


def _about_sheet(wb: Workbook, tables: dict[str, pd.DataFrame], rates: pd.Series) -> None:
    ws = wb["About"]
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    _heading(ws, "About this workbook")
    records = tables["records"]
    used = records[records["round_id"] != ""]
    lines = [
        ("What it is", "Every Indian startup funding round the tracker found since "
                       f"{START_DATE:%d %B %Y}, one row per round, with the reports that support it."),
        ("Sources", "Entrackr and The Economic Times (articles), Inc42 (weekly deal tables), VCCircle (headlines). "
                    "Only facts and links are stored, never article text."),
        ("Reports read", "; ".join(f"{SOURCE_LABELS[s]} {int((used['source'] == s).sum()):,}"
                                   for s in SOURCE_LABELS) + " matched to a round."),
        ("Exchange rate", "Rupee amounts are converted at the ECB reference rate for the announcement date "
                          f"(latest: {rates.iloc[-1]:.2f} INR per USD on {rates.index[-1]:%d %b %Y})."),
        ("PivotTables", "By Sector, By Stage and By City are PivotTables on the Rounds table. They refresh when the "
                        "file opens. After editing data: PivotTable Analyze > Refresh."),
        ("Checks", f"{len(CHECKS)} automated checks run on every build. Results are on the Checks and Exceptions sheets."),
        ("Rules", "The Source rules sheet says which source wins for each field, and why."),
        ("Rebuild", "python -m funding_tracker update   (see README.md in the repository)"),
        ("Repository", "https://github.com/pkjx29/funding-tracker"),
    ]
    for row, (label, text) in enumerate(lines, 3):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color=INK)
        ws.cell(row=row, column=2, value=text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=1).alignment = Alignment(vertical="top")


# ---------------------------------------------------------- cached formulas
def add_cached_values(path: Path, cached: dict[str, dict[str, object]]) -> None:
    """Store each formula's result next to the formula.

    openpyxl writes formulas without their results. Excel recalculates on
    opening, but previews (Finder Quick Look, mail clients, phones) show only
    stored results, so without this step the dashboard would look empty there.
    """
    temp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        workbook = source.read("xl/workbook.xml").decode("utf-8")
        relations = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        targets = dict(re.findall(r'Target="([^"]+)" Id="([^"]+)"', relations))
        targets = {rel_id: target_path.lstrip("/") for target_path, rel_id in targets.items()}
        sheet_files = {}
        for match in re.finditer(r"<sheet [^>]*>", workbook):
            name = re.search(r'name="([^"]*)"', match.group(0)).group(1)
            rel_id = re.search(r'r:id="([^"]*)"', match.group(0)).group(1)
            file_name = targets[rel_id]
            sheet_files[file_name if file_name.startswith("xl/") else f"xl/{file_name}"] = name
        for item in source.infolist():
            data = source.read(item.filename)
            sheet = sheet_files.get(item.filename)
            if sheet in cached:
                xml = data.decode("utf-8")
                for ref, value in cached[sheet].items():
                    if isinstance(value, str):
                        shown, cell_type = escape(value), ' t="str"'
                    elif isinstance(value, date):              # dates are stored as serial numbers
                        shown, cell_type = str((value - date(1899, 12, 30)).days), ""
                    else:
                        shown, cell_type = repr(value), ""
                    pattern = re.compile(rf'(<c r="{ref}"(?: [a-z]+="[^"]*")*)(><f>.*?</f>)<v></v>')
                    xml, count = pattern.subn(
                        lambda m: m.group(1) + cell_type + m.group(2) + "<v>" + shown + "</v>", xml, count=1)
                    if count != 1:
                        raise ValueError(f"no formula found at {sheet}!{ref}")
                data = xml.encode("utf-8")
            target.writestr(item, data)
    shutil.move(temp, path)


# --------------------------------------------------------------------- main
SHEETS = ["Dashboard", "Rounds", "By Sector", "By Stage", "By City", "Checks", "Exceptions", "Source rules",
          "Reports", "Coverage", "Chart data", "About"]


def write_workbook(tables: dict[str, pd.DataFrame], path: Path, rates: pd.Series) -> None:
    rounds, records = tables["rounds"], tables["records"]
    frame = rounds_frame(rounds)
    orders = display_orders(frame)
    cached: dict[str, dict[str, object]] = {}

    wb = Workbook()
    wb.active.title = SHEETS[0]
    for name in SHEETS[1:]:
        wb.create_sheet(name)

    # ---- Rounds: the table everything else reads from
    ws = wb["Rounds"]
    last = _write_table(ws, frame, "Rounds", [w for *_, w, _ in ROUND_COLUMNS], [f for *_, f in ROUND_COLUMNS])
    link_col = HEADERS.index("First source link") + 1
    for row in range(2, last + 1):
        cell = ws.cell(row=row, column=link_col)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(color="1C5CAB", underline="single")
    ws.freeze_panes = "C2"

    # ---- PivotTables
    cache = _pivot_cache(frame, orders)
    for sheet, field, pivot_name in (("By Sector", "Sector", "PivotSector"), ("By Stage", "Stage", "PivotStage"),
                                     ("By City", "City", "PivotCity")):
        ws = wb[sheet]
        _heading(ws, f"Rounds by {field.lower()}",
                 "PivotTable on the Rounds table. It refreshes when the file opens; drag fields to re-cut it.")
        _add_pivot(ws, frame, cache, field, orders[field], pivot_name)

    # ---- reports behind the rounds
    used = records[records["round_id"] != ""].sort_values(["round_id", "published_date", "record_id"])
    reports = pd.DataFrame({
        "Record ID": used["record_id"], "Round ID": used["round_id"],
        "Source": used["source"].map(SOURCE_LABELS), "Published": pd.to_datetime(used["published_date"]).dt.date,
        "Company in report": used["company_raw"], "Amount as reported": used["amount_text"],
        "INR cr": used["amount_inr_cr"], "USD mn": used["amount_usd_mn"], "Stage": used["stage"],
        "Lead investors": used["lead_investors"], "Link": used["url"],
    })
    _plain_sheet(wb, "Reports", "Source reports matched to a round",
                 "One row per article, table row or headline. Amounts are exactly as each source gave them.",
                 reports, "Reports", [17, 24, 20, 13, 30, 18, 11, 11, 14, 36, 60],
                 [None, None, None, DATE_FORMAT, None, None, "#,##0.00", "#,##0.00", None, None, None])

    _chart_data(wb, frame, orders, cached)
    _dashboard(wb, frame, tables, len(reports), cached)
    _checks_sheet(wb, tables["checks_summary"])

    exceptions = tables["exceptions"].rename(columns={
        "check_id": "Check ID", "check": "Check", "severity": "Severity", "round_id": "Round ID",
        "company": "Company", "announced_date": "Announced", "detail": "Detail"})
    exceptions["Announced"] = pd.to_datetime(exceptions["Announced"]).dt.date
    _plain_sheet(wb, "Exceptions", "Exceptions: every flag raised by the checks",
                 "Filter by Check ID or Severity. The Flags column on the Rounds sheet lists the same IDs per round.",
                 exceptions, "Exceptions", [10, 36, 11, 24, 28, 13, 110],
                 [None, None, None, None, None, DATE_FORMAT, None])

    rules = pd.DataFrame(SOURCE_RULES, columns=["Field", "Source-of-truth rule", "Why"])
    _plain_sheet(wb, "Source rules", "Source of truth, field by field",
                 "When sources disagree, these rules decide what goes on the Rounds sheet. The disagreement is still flagged.",
                 rules, "SourceRules", [26, 90, 80], [None, None, None])
    for row in wb["Source rules"].iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    coverage = tables["coverage"].rename(columns={
        "week_start": "Week start", "week_end": "Week end", "entrackr_deals": "Entrackr: funded startups",
        "tracked_rounds": "Tracker: rounds", "difference": "Difference",
        "entrackr_usd_mn": "Entrackr: USD mn", "tracked_usd_mn": "Tracker: USD mn"})
    for column_name in ("Week start", "Week end"):
        coverage[column_name] = pd.to_datetime(coverage[column_name]).dt.date
    _plain_sheet(wb, "Coverage", "Coverage check against Entrackr's weekly report",
                 "A control total: the number of funded startups Entrackr counted each week, beside the rounds tracked here.",
                 coverage, "Coverage", [13, 13, 26, 16, 12, 18, 18],
                 [DATE_FORMAT, DATE_FORMAT, "#,##0", "#,##0", "#,##0", "#,##0.00", "#,##0.00"])

    _about_sheet(wb, tables, rates)
    wb["Chart data"].sheet_properties.tabColor = FAINT
    cached.pop("_extent")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    add_cached_values(path, cached)
